import openpyxl
import json
import os

wb = openpyxl.load_workbook('backend/seed/curriculum_dataset.xlsx')
sheet = wb.active
headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
col_map = {str(h).strip(): idx for idx, h in enumerate(headers) if h is not None}

def parse_semester(val, code):
    s = str(val).strip()
    if s.isdigit():
        return int(s)
    if 'Final' in s:
        return 10
    if 'Elective' in s:
        return 9
    if code == 'ART 204':
        return 2
    return 8

courses = []
for row in sheet.iter_rows(min_row=2, values_only=True):
    if not row or not any(row):
        continue
    code = row[col_map.get('Course Code')]
    name = row[col_map.get('Course Name')]
    if not code or not name:
        continue
    credits = row[col_map.get('Credit Hours')]
    try:
        credits = float(credits) if credits is not None else 3.0
    except:
        credits = 3.0
    code_str = str(code).strip().upper()
    semester = parse_semester(row[col_map.get('Semester')], code_str)
    desc = row[col_map.get('Description')] or ''
    topics = row[col_map.get('Topics / Skills')] or ''
    assignments = row[col_map.get('Projects / Assignments')] or ''
    co1 = row[col_map.get('CO1')] or ''
    co2 = row[col_map.get('CO2')] or ''
    co3 = row[col_map.get('CO3')] or ''
    co4 = row[col_map.get('CO4')] or ''
    co5 = row[col_map.get('CO5')] or ''

    courses.append({
        'code': code_str,
        'title': str(name).strip(),
        'credits': credits,
        'semester': semester,
        'description': str(desc).strip(),
        'topics': str(topics).strip(),
        'assignments': str(assignments).strip(),
        'co1': str(co1).strip(),
        'co2': str(co2).strip(),
        'co3': str(co3).strip(),
        'co4': str(co4).strip(),
        'co5': str(co5).strip(),
    })

os.makedirs('supabase', exist_ok=True)
with open('supabase/courses_dataset.json', 'w', encoding='utf-8') as f:
    json.dump(courses, f, indent=2)

# Generate seed.sql
with open('supabase/seed.sql', 'w', encoding='utf-8') as f:
    f.write('-- Master Curriculum Seed (82 IUBAT CSE Courses)\n')
    f.write('INSERT INTO public.courses (course_code, course_name, credit_hours, semester, description, topics, assignments, co1, co2, co3, co4, co5)\nVALUES\n')
    lines = []
    for c in courses:
        def escape_sql(val):
            return "'" + str(val).replace("'", "''") + "'"
        line = f"({escape_sql(c['code'])}, {escape_sql(c['title'])}, {c['credits']}, {c['semester']}, {escape_sql(c['description'])}, {escape_sql(c['topics'])}, {escape_sql(c['assignments'])}, {escape_sql(c['co1'])}, {escape_sql(c['co2'])}, {escape_sql(c['co3'])}, {escape_sql(c['co4'])}, {escape_sql(c['co5'])})"
        lines.append(line)
    f.write(',\n'.join(lines))
    f.write('\nON CONFLICT (course_code) DO UPDATE SET\n')
    f.write('  course_name = EXCLUDED.course_name,\n')
    f.write('  credit_hours = EXCLUDED.credit_hours,\n')
    f.write('  semester = EXCLUDED.semester,\n')
    f.write('  description = EXCLUDED.description,\n')
    f.write('  topics = EXCLUDED.topics,\n')
    f.write('  assignments = EXCLUDED.assignments,\n')
    f.write('  co1 = EXCLUDED.co1, co2 = EXCLUDED.co2, co3 = EXCLUDED.co3, co4 = EXCLUDED.co4, co5 = EXCLUDED.co5;\n')

# Also write TypeScript module for frontend
with open('src/data/curriculum-dataset.ts', 'w', encoding='utf-8') as f:
    f.write('/**\n * Real IUBAT CSE Curriculum Dataset (82 courses extracted from official syllabi)\n */\n')
    f.write('export interface CurriculumCourseData {\n')
    f.write('  code: string;\n')
    f.write('  title: string;\n')
    f.write('  credits: number;\n')
    f.write('  semester: number;\n')
    f.write('  description: string;\n')
    f.write('  topics?: string;\n')
    f.write('  assignments?: string;\n')
    f.write('  co1?: string;\n')
    f.write('  co2?: string;\n')
    f.write('  co3?: string;\n')
    f.write('  co4?: string;\n')
    f.write('  co5?: string;\n')
    f.write('}\n\n')
    f.write('export const masterCurriculumCourses: CurriculumCourseData[] = ')
    json.dump(courses, f, indent=2)
    f.write(';\n')

print(f'Successfully extracted and generated files for {len(courses)} courses!')
