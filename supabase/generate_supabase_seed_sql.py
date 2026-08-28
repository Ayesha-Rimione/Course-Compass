import os
import sys
import openpyxl

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "backend")))

def escape_sql(val):
    if val is None:
        return "NULL"
    return "'" + str(val).replace("'", "''") + "'"

def generate_seed_sql():
    excel_path = os.path.join(os.path.dirname(__file__), "..", "backend", "seed", "curriculum_dataset.xlsx")
    wb = openpyxl.load_workbook(excel_path)
    sheet = wb.active

    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    headers_clean = [str(h).strip() if h is not None else "" for h in headers]
    col_map = {h: idx for idx, h in enumerate(headers_clean)}

    out_file = os.path.join(os.path.dirname(__file__), "seed.sql")

    with open(out_file, "w", encoding="utf-8") as f:
        f.write("-- ==============================================================================\n")
        f.write("-- CourseCompass — Complete Supabase PostgreSQL Seed Data\n")
        f.write("-- ==============================================================================\n\n")

        # 1. Departments & Programs
        f.write("-- 1. Departments & Programs\n")
        f.write("INSERT INTO public.departments (id, code, name)\n")
        f.write("VALUES ('a0000000-0000-0000-0000-000000000001', 'CSE', 'Department of Computer Science & Engineering')\n")
        f.write("ON CONFLICT (code) DO NOTHING;\n\n")

        f.write("INSERT INTO public.programs (id, department_id, code, name, degree_type, total_credits)\n")
        f.write("VALUES ('b0000000-0000-0000-0000-000000000001', 'a0000000-0000-0000-0000-000000000001', 'BCSE', 'Bachelor of Science in Computer Science & Engineering', 'Bachelor of Science', 140)\n")
        f.write("ON CONFLICT (code) DO NOTHING;\n\n")

        f.write("INSERT INTO public.curriculum_versions (id, program_id, name, version_code, effective_from, status)\n")
        f.write("VALUES ('c0000000-0000-0000-0000-000000000001', 'b0000000-0000-0000-0000-000000000001', 'IUBAT CSE Curriculum (B231)', 'B231', '2023-01', 'ACTIVE')\n")
        f.write("ON CONFLICT DO NOTHING;\n\n")

        # 2. All 82 Courses
        f.write("-- 2. Master Curriculum (82 IUBAT CSE Courses)\n")
        f.write("INSERT INTO public.courses (course_code, course_name, credit_hours, semester, description)\nVALUES\n")
        course_lines = []
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if not row or all(v is None for v in row):
                continue
            raw_code = row[col_map["Course Code"]] if "Course Code" in col_map else None
            raw_name = row[col_map["Course Name"]] if "Course Name" in col_map else None
            if not raw_code or not raw_name or str(raw_code).strip() == "":
                continue

            course_code = str(raw_code).strip().upper()
            course_name = str(raw_name).replace("\ufffd", " ").strip()
            raw_credits = row[col_map["Credit Hours"]] if "Credit Hours" in col_map else 3.0
            try:
                credit_hours = float(raw_credits) if raw_credits is not None else 3.0
            except:
                credit_hours = 3.0

            raw_sem = row[col_map["Semester"]] if "Semester" in col_map else 1
            if str(raw_sem).strip().isdigit():
                semester = int(str(raw_sem).strip())
            elif "Final" in str(raw_sem):
                semester = 10
            elif "Elective" in str(raw_sem):
                semester = 9
            elif course_code == "ART 204":
                semester = 2
            else:
                semester = 8

            description = str(row[col_map["Description"]]).replace("\ufffd", " ").strip() if "Description" in col_map and row[col_map["Description"]] else ""

            course_lines.append(f"({escape_sql(course_code)}, {escape_sql(course_name)}, {credit_hours}, {semester}, {escape_sql(description)})")

        f.write(",\n".join(course_lines))
        f.write("\nON CONFLICT (course_code) DO UPDATE SET\n")
        f.write("  course_name = EXCLUDED.course_name,\n")
        f.write("  credit_hours = EXCLUDED.credit_hours,\n")
        f.write("  semester = EXCLUDED.semester,\n")
        f.write("  description = EXCLUDED.description;\n\n")

        # 3. 29 Skills
        f.write("-- 3. Skills Catalog (29 Skills)\n")
        f.write("INSERT INTO public.skills (name, category, description)\nVALUES\n")
        skills_data = [
            ("Python", "Programming", "Versatile high-level programming language for web, scripting, and AI/ML."),
            ("JavaScript", "Programming", "Core programming language of the modern web."),
            ("TypeScript", "Programming", "Typed superset of JavaScript for scalable applications."),
            ("C++", "Programming", "High-performance systems programming language."),
            ("C", "Programming", "Foundational procedural systems language."),
            ("Java", "Programming", "Object-oriented language for enterprise systems and Android."),
            ("SQL", "Database", "Relational database querying and management."),
            ("PostgreSQL", "Database", "Advanced open-source relational database system."),
            ("MongoDB", "Database", "Document-oriented NoSQL database."),
            ("React", "Frontend", "Declarative UI component library."),
            ("Node.js", "Backend", "Asynchronous event-driven JavaScript runtime."),
            ("FastAPI", "Backend", "High-performance Python asynchronous REST API framework."),
            ("Django", "Backend", "Batteries-included Python web framework."),
            ("Docker", "DevOps", "Container platform for packaging and deploying applications."),
            ("Kubernetes", "DevOps", "Container orchestration system for automating deployment."),
            ("CI/CD", "DevOps", "Continuous Integration and Continuous Deployment pipelines."),
            ("Linux", "Tools & DevOps", "Unix-like open source operating system."),
            ("Git", "Tools & DevOps", "Distributed version control system."),
            ("Machine Learning", "AI & Data Science", "Algorithms and statistical models that learn from data."),
            ("Deep Learning", "AI & Data Science", "Neural networks and deep representations."),
            ("Data Structures", "Computer Science", "Fundamental data organization (arrays, trees, graphs)."),
            ("Algorithms", "Computer Science", "Step-by-step problem solving and computational efficiency."),
            ("System Design", "Architecture", "Designing scalable distributed software architectures."),
            ("Cybersecurity", "Security", "Protecting systems, networks, and data from attacks."),
            ("Cloud Architecture", "Cloud", "Designing resilient cloud infrastructure on AWS/GCP/Azure."),
            ("HTML/CSS", "Web Development", "Fundamental building blocks of web pages."),
            ("Next.js", "Web Development", "React production framework for full-stack applications."),
            ("Data Analysis", "AI & Data Science", "Inspecting, cleansing, and transforming data for decisions."),
            ("Problem Solving", "Soft Skills", "Analytical mindset for diagnosing and resolving complex challenges.")
        ]
        skill_lines = [f"({escape_sql(n)}, {escape_sql(c)}, {escape_sql(d)})" for n, c, d in skills_data]
        f.write(",\n".join(skill_lines))
        f.write("\nON CONFLICT (name) DO UPDATE SET category = EXCLUDED.category, description = EXCLUDED.description;\n\n")

        # 4. 10 Career Roles
        f.write("-- 4. Career Roles (10 Roles)\n")
        f.write("INSERT INTO public.career_roles (name, description, category)\nVALUES\n")
        roles_data = [
            ("Full-Stack Software Engineer", "Build end-to-end web applications across frontend, backend, and databases.", "Software Engineering"),
            ("Backend Engineer", "Design, build, and optimize backend microservices, APIs, and data layers.", "Software Engineering"),
            ("Frontend Engineer", "Create polished, performant, and accessible user interfaces for web applications.", "Frontend Engineering"),
            ("AI/ML Engineer", "Develop, train, and deploy machine learning models and intelligent services into production.", "Artificial Intelligence"),
            ("DevOps & Cloud Engineer", "Automate infrastructure, build CI/CD pipelines, and maintain highly available cloud systems.", "DevOps & Cloud"),
            ("Cybersecurity Analyst", "Protect networks, systems, and sensitive applications from digital vulnerabilities.", "Security"),
            ("Data Scientist", "Extract actionable insights, statistical models, and predictions from massive data.", "Data Science"),
            ("Data Analyst", "Transform raw data into dashboards, business metrics, and strategic reporting.", "Data Analytics"),
            ("Mobile App Developer", "Build native and cross-platform mobile apps for iOS and Android platforms.", "Mobile Engineering"),
            ("Systems Engineer", "Build low-level operating system modules, compilers, and performance-critical systems.", "Systems Engineering")
        ]
        role_lines = [f"({escape_sql(n)}, {escape_sql(d)}, {escape_sql(c)})" for n, d, c in roles_data]
        f.write(",\n".join(role_lines))
        f.write("\nON CONFLICT (name) DO UPDATE SET description = EXCLUDED.description, category = EXCLUDED.category;\n\n")

        print("Supabase seed.sql generated successfully!")

if __name__ == "__main__":
    generate_seed_sql()
